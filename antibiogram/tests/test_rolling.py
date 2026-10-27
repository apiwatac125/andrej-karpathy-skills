"""เทสต์ย่อสำหรับ ab.rolling + export.to_paired_excel (รันได้ด้วย `python tests/test_rolling.py`).

สร้างข้อมูลสังเคราะห์ที่มีวันที่ แล้วตรวจว่า:
- หน้าต่างรายปี/rolling กรอง collect_datetime ถูกช่วง
- compare() ให้ Δ = rolling − annual ต่อคู่ (เชื้อ x ยา)
- to_paired_excel() ได้ไฟล์ 3 แผ่น Annual / Rolling / Compare
"""
import io
import random
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from ab import config as cfg, export, rolling  # noqa: E402

AB_COLS = ["CAZ", "CRO", "FEP", "MEM", "CIP"]


def _make(org, year, n, s_prob, rng):
    rows = []
    for i in range(n):
        dt = pd.Timestamp(year=year, month=rng.randint(1, 12), day=rng.randint(1, 28))
        row = {"hn": f"{org[:2]}{year}{i:04d}", "organism": org, "specimen": "Blood",
               "ward": "ICU", "infection_origin": "CAI", "collect_datetime": dt}
        for ab in AB_COLS:
            row[ab] = "S" if rng.random() < s_prob.get(ab, 0.8) else "R"
        rows.append(row)
    return rows


def _synthetic():
    rng = random.Random(1)
    rows = []
    for org in ["Escherichia coli", "Klebsiella pneumoniae"]:
        rows += _make(org, 2024, 80, {"CAZ": .90, "CRO": .85, "FEP": .88, "MEM": .99, "CIP": .70}, rng)
        rows += _make(org, 2025, 80, {"CAZ": .55, "CRO": .60, "FEP": .70, "MEM": .97, "CIP": .50}, rng)
    return pd.DataFrame(rows)


def main():
    conditions = cfg.load_conditions()
    std = _synthetic()

    ann = rolling.annual(std, AB_COLS, conditions, year=2025, origins=["CAI", "HAI"])
    rol = rolling.rolling(std, AB_COLS, conditions, end="2025-06-30", months=12, origins=["CAI", "HAI"])

    assert ann.window["kind"] == "annual" and ann.window["year"] == 2025
    assert rol.window["kind"] == "rolling" and rol.window["months"] == 12
    assert rol.window["start"].date().isoformat() == "2024-06-30"
    assert rol.window["end"].date().isoformat() == "2025-06-30"
    # annual นับเฉพาะปี 2025 (80 ตัว/เชื้อ หลัง dedup hn ไม่ซ้ำ)
    assert ann.totals.get("Escherichia coli") == 80

    cmp = rolling.compare(ann, rol)
    assert not cmp.empty
    assert set(["organism", "antibiotic", "pct_annual", "pct_rolling", "delta"]).issubset(cmp.columns)
    assert (cmp["delta"] == cmp["pct_rolling"] - cmp["pct_annual"]).all()

    data = export.to_paired_excel(ann, rol, scope="All specimen & ward - ALL")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Annual", "Rolling", "Compare"], wb.sheetnames

    print(f"OK — annual={ann.window['label']} · rolling={rol.window['label']} · "
          f"compare rows={len(cmp)} · excel {len(data)} bytes")


if __name__ == "__main__":
    main()
