"""ส่งออกผล antibiogram เป็น Excel รูปแบบเดียวกับไฟล์ publish ของโรงพยาบาล.

หน้าตา:
- แถวหัวเรื่อง + ขอบเขต
- หัวตารางเขียว (006600 ตัวอักษรขาว): แถว class (merge) + แถวชื่อยา
- แต่ละเชื้อ 2 แถว: บน = %S, ล่าง = จำนวน N (ระบายสีตามแบนด์ %S)
- ช่อง intrinsic = "R"
- คอลัมน์ Organism Name, Quality (จำนวนรวม)
"""

from __future__ import annotations

import io
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CONFIG_DIR = Path(__file__).resolve().parent.parent / "config"
CLASSES_PATH = CONFIG_DIR / "antibiotic_classes.csv"

# สีตามแบนด์ %S (ตรงกับไฟล์ publish)
GREEN, YELLOW, ORANGE, RED, GREY = "FF92D050", "FFFFFF00", "FFFFC000", "FFFF0000", "FF7F7F7F"
HEAD = "FF006600"
THIN = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _band_fill(pct: float) -> str:
    if pct >= 90:
        return GREEN
    if pct >= 80:
        return YELLOW
    if pct >= 70:
        return ORANGE
    if pct >= 51:
        return RED
    return GREY


def _load_classes(path: Path = CLASSES_PATH) -> dict[str, tuple[str, int]]:
    df = pd.read_csv(path, dtype={"standard_name": str, "class": str, "order": int})
    return {r["standard_name"].strip(): (r["class"].strip(), int(r["order"]))
            for _, r in df.iterrows()}


def to_publish_excel(
    long_form: pd.DataFrame,
    drug_name: dict[str, str],
    organism_totals: dict[str, int],
    title: str = "Percentage of susceptible bacteria",
    scope: str = "All specimen & ward",
    date_range: str = "",
    organism_order: list[str] | None = None,
    classes_path: Path = CLASSES_PATH,
) -> bytes:
    classes = _load_classes(classes_path)

    # คอลัมน์ยาที่จะแสดง = ยาที่มีในข้อมูล และรู้จักใน antibiotic_classes เรียงตาม order
    present_std = {}
    for ab in long_form["antibiotic"].unique():
        std = drug_name.get(ab, ab)
        if std in classes and std not in present_std:
            present_std[std] = ab  # std -> ab_col ตัวแรกที่เจอ
    ordered = sorted(present_std, key=lambda s: classes[s][1])
    if not ordered:
        ordered = sorted(present_std)

    # cell lookup: (organism, std_drug) -> (value, n)
    cell = {}
    for _, r in long_form.iterrows():
        std = drug_name.get(r["antibiotic"], r["antibiotic"])
        if bool(r.get("intrinsic")):
            val = "R"
        elif r["reportable"]:
            val = r["pct_susceptible"]
        else:
            val = None
        cell[(r["organism"], std)] = (val, int(r["n_tested"]))

    organisms = organism_order or sorted(
        organism_totals, key=lambda o: organism_totals.get(o, 0), reverse=True
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Antibiogram"
    ncol = 2 + len(ordered)  # A=organism, B=quality, then drugs

    # หัวเรื่อง
    ws.cell(row=1, column=2, value=title)
    ws.cell(row=2, column=2, value=date_range)
    ws.cell(row=3, column=2, value=scope)
    r0 = 4  # แถวหัว class

    head_font = Font(bold=True, color="FFFFFFFF")
    head_fill = PatternFill("solid", fgColor=HEAD)

    def style_head(c):
        c.fill = head_fill; c.font = head_font; c.alignment = CENTER; c.border = BORDER

    # Organism Name / Quality (merge 2 แถว)
    ws.merge_cells(start_row=r0, start_column=1, end_row=r0 + 1, end_column=1)
    ws.merge_cells(start_row=r0, start_column=2, end_row=r0 + 1, end_column=2)
    c = ws.cell(row=r0, column=1, value="Organism Name"); style_head(c)
    c = ws.cell(row=r0, column=2, value="Quality"); style_head(c)

    # class header (merge ตามจำนวนยาในแต่ละ class ที่ปรากฏ) + drug row
    col = 3
    i = 0
    while i < len(ordered):
        cls = classes[ordered[i]][0]
        j = i
        while j < len(ordered) and classes[ordered[j]][0] == cls:
            j += 1
        span = j - i
        if span > 1:
            ws.merge_cells(start_row=r0, start_column=col, end_row=r0, end_column=col + span - 1)
        c = ws.cell(row=r0, column=col, value=cls); style_head(c)
        for k in range(span):
            dc = ws.cell(row=r0 + 1, column=col + k, value=ordered[i + k]); style_head(dc)
        col += span
        i = j

    # ข้อมูลเชื้อ (2 แถวต่อเชื้อ)
    row = r0 + 2
    for org in organisms:
        top, bot = row, row + 1
        ws.merge_cells(start_row=top, start_column=1, end_row=bot, end_column=1)
        ws.merge_cells(start_row=top, start_column=2, end_row=bot, end_column=2)
        oc = ws.cell(row=top, column=1, value=org)
        oc.alignment = Alignment(vertical="center", wrap_text=True); oc.border = BORDER
        qc = ws.cell(row=top, column=2, value=organism_totals.get(org))
        qc.alignment = CENTER; qc.border = BORDER; qc.font = Font(bold=True)
        for idx, std in enumerate(ordered):
            ccol = 3 + idx
            val, n = cell.get((org, std), (None, 0))
            tcell = ws.cell(row=top, column=ccol)
            bcell = ws.cell(row=bot, column=ccol)
            tcell.alignment = CENTER; bcell.alignment = CENTER
            tcell.border = BORDER; bcell.border = BORDER
            if val == "R":
                tcell.value = "R"
            elif val is not None:
                tcell.value = val
                bcell.value = n
                bcell.fill = PatternFill("solid", fgColor=_band_fill(val))
        row += 2

    # footnotes
    ws.cell(row=row + 1, column=1, value="( ) = จำนวนเชื้อที่ทดสอบ")
    ws.cell(row=row + 2, column=1, value="ยา Colistin เป็นการคำนวณจาก %I เท่านั้น")

    # ความกว้างคอลัมน์
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 8
    for idx in range(len(ordered)):
        ws.column_dimensions[get_column_letter(3 + idx)].width = 6

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_excel_bytes(matrix: pd.DataFrame, long_form: pd.DataFrame) -> bytes:
    """ส่งออกแบบเรียบง่าย (สรุป + รายละเอียด) — ใช้เมื่อไม่ต้องการรูปแบบ publish."""
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        matrix.to_excel(writer, sheet_name="Antibiogram")
        long_form.to_excel(writer, sheet_name="Details", index=False)
    return buffer.getvalue()
